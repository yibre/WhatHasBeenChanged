"""
diff_engine.py — 두 엑셀 파일의 변경점을 git처럼 추적하는 엔진.

브라우저(Pyodide)와 일반 CPython에서 동일하게 동작한다.
파일 바이트를 받아 시트/열/행/셀 단위의 추가·삭제·수정을 계산하고
JSON 문자열로 돌려준다.

알고리즘
  1. 시트 이름으로 매칭 (한쪽에만 있으면 시트 추가/삭제)
  2. 열 정렬  : 각 열의 전체 값을 키로 만들어 difflib으로 LCS 정렬
  3. 행 정렬  : 매칭된 열의 값만으로 행 키를 만들어 difflib으로 LCS 정렬
                 (replace 구간은 유사도 기반으로 짝지어 '수정된 행'으로 판정)
  4. 셀 비교  : 짝지어진 행/열 교차점의 값을 비교

CLI 사용:
    python diff_engine.py 이전.xlsx 이후.xlsx
    python diff_engine.py 이전.xlsx 이후.xlsx --json out.json --report 리포트.xlsx
"""

from __future__ import annotations

import datetime as _dt
import io
import json
from difflib import SequenceMatcher

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

__all__ = ["compare", "report_xlsx", "DEFAULTS"]

# 셀/행/열 상태 코드 (JSON 크기를 줄이려고 숫자로 쓴다)
SAME, MOD, ADD, DEL, NONE = 0, 1, 2, 3, 4

DEFAULTS = {
    "compare": "value",      # "value" = 계산된 값 기준, "formula" = 수식 기준
    "ignore_case": False,    # 대소문자 무시
    "ignore_space": False,   # 공백 무시
    "trim": True,            # 앞뒤 공백 무시
    "similarity": 0.4,       # 이 정도 이상 닮은 행은 '삭제+추가'가 아니라 '수정'으로 본다
    "view": "auto",          # "auto" | "all" | "changes"
    "context": 3,            # 변경 주변 몇 행까지 같이 보여줄지
}

MAX_ROWS = 30000            # 이보다 큰 시트는 잘라서 비교 (브라우저 보호)
MAX_COLS = 2000
MAX_PAIR_BLOCK = 400        # 유사도 짝짓기를 시도할 최대 블록 크기
MAX_CHANGE_LIST = 20000     # 변경 목록 최대 건수

_CACHE: dict = {"sig": None, "a": None, "b": None}
_LAST: dict | None = None


# ---------------------------------------------------------------- 값 정규화

def _fmt(v) -> str:
    """셀 값을 화면/비교용 문자열로 바꾼다."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, _dt.datetime):
        if v.hour == v.minute == v.second == 0 and v.microsecond == 0:
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, _dt.date):
        return v.isoformat()
    if isinstance(v, _dt.time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, _dt.timedelta):
        return str(v)
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):
            return str(v)
        if abs(v) < 1e15 and float(v).is_integer():
            return str(int(v))
        return f"{v:.10g}"
    if isinstance(v, int):
        return str(v)
    return str(v)


def _mkkey(opts):
    trim, low, nospace = opts["trim"], opts["ignore_case"], opts["ignore_space"]

    def key(t: str) -> str:
        if nospace:
            t = "".join(t.split())
        elif trim:
            t = t.strip()
        if low:
            t = t.casefold()
        return t

    return key


# ---------------------------------------------------------------- 파일 읽기

def _read_sheet(ws) -> list[list]:
    rows, width = [], 0
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i >= MAX_ROWS:
            break
        row = list(r[:MAX_COLS])
        width = max(width, len(row))
        rows.append(row)
    for row in rows:
        if len(row) < width:
            row.extend([None] * (width - len(row)))
    return rows


def _to_bytes(data) -> bytes:
    """JS Uint8Array / memoryview / bytes 무엇이 와도 bytes로."""
    if isinstance(data, (bytes, bytearray)):
        return bytes(data)
    to_py = getattr(data, "to_py", None)
    if to_py is not None:
        data = to_py()
    return bytes(data)


def _parse(data: bytes, mode: str) -> dict:
    """워크북을 {시트이름: 값 그리드} 로 읽는다."""
    wb = load_workbook(io.BytesIO(data), data_only=False, read_only=True, keep_links=False)
    order = list(wb.sheetnames)
    grids = {name: _read_sheet(wb[name]) for name in order}
    wb.close()

    if mode == "value":
        # data_only=True 는 엑셀이 저장해 둔 계산 결과를 준다.
        # 계산 결과가 없는 셀(파이썬으로 만든 파일 등)은 수식 그리드로 메운다.
        wbv = load_workbook(io.BytesIO(data), data_only=True, read_only=True, keep_links=False)
        for name in wbv.sheetnames:
            gv = _read_sheet(wbv[name])
            gf = grids.get(name, [])
            merged = []
            for i in range(max(len(gf), len(gv))):
                rf = gf[i] if i < len(gf) else []
                rv = gv[i] if i < len(gv) else []
                w = max(len(rf), len(rv))
                merged.append([
                    rv[j] if j < len(rv) and rv[j] is not None
                    else (rf[j] if j < len(rf) else None)
                    for j in range(w)
                ])
            grids[name] = merged
        wbv.close()

    return {"order": order, "grids": grids}


def _prep(grid: list[list], keyfn) -> tuple[list[list[str]], list[list[str]]]:
    """직사각형으로 맞추고, 뒤쪽 빈 행/열을 잘라내고, 표시용·비교용 그리드를 만든다."""
    texts = [[_fmt(v) for v in row] for row in grid]
    w = max((len(r) for r in texts), default=0)
    for r in texts:
        if len(r) < w:
            r.extend([""] * (w - len(r)))

    while texts and all(c == "" for c in texts[-1]):
        texts.pop()
    while w > 0 and texts and all(r[w - 1] == "" for r in texts):
        for r in texts:
            r.pop()
        w -= 1

    keys = [[keyfn(c) for c in row] for row in texts]
    return texts, keys


# ---------------------------------------------------------------- 정렬(매칭)

def _opcode_pairs(a: list, b: list, pair_block):
    """두 수열을 정렬해 (a인덱스, b인덱스, 상태) 목록으로 만든다."""
    out = []
    for tag, i1, i2, j1, j2 in SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if tag == "equal":
            out += [(i, j, SAME) for i, j in zip(range(i1, i2), range(j1, j2))]
        elif tag == "delete":
            out += [(i, None, DEL) for i in range(i1, i2)]
        elif tag == "insert":
            out += [(None, j, ADD) for j in range(j1, j2)]
        else:
            out += pair_block(list(range(i1, i2)), list(range(j1, j2)))
    return out


def _positional(ai: list[int], bj: list[int]):
    """replace 구간을 위치 순서대로 짝짓는다 (열에 사용)."""
    out = []
    n = min(len(ai), len(bj))
    out += [(ai[k], bj[k], MOD) for k in range(n)]
    out += [(i, None, DEL) for i in ai[n:]]
    out += [(None, j, ADD) for j in bj[n:]]
    return out


def _similar(ka: tuple, kb: tuple) -> float:
    if not ka and not kb:
        return 1.0
    n = max(len(ka), len(kb))
    same = sum(1 for x, y in zip(ka, kb) if x == y)
    return same / n if n else 1.0


def _make_row_pairer(keys_a, keys_b, threshold):
    """replace 구간에서 '닮은 행'끼리 순서를 지키며 짝짓는다 (행에 사용)."""

    def pair(ai, bj):
        if not ai:
            return [(None, j, ADD) for j in bj]
        if not bj:
            return [(i, None, DEL) for i in ai]
        if len(ai) > MAX_PAIR_BLOCK or len(bj) > MAX_PAIR_BLOCK:
            return _positional(ai, bj)

        # 같은 개수의 행이 통째로 바뀐 구간은 위치끼리 대응시킨다.
        # (열이 몇 개 없는 표에서 셀 하나 바뀐 행이 '삭제+추가'로 갈라지는 걸 막는다)
        if len(ai) == len(bj):
            avg = sum(_similar(keys_a[i], keys_b[j]) for i, j in zip(ai, bj)) / len(ai)
            width = max((len(keys_a[i]) for i in ai), default=0)
            if width <= 3 or avg >= min(threshold, 0.25):
                return _positional(ai, bj)

        n, m = len(ai), len(bj)
        # 순서를 유지하며 유사도 합을 최대화하는 DP
        dp = [[0.0] * (m + 1) for _ in range(n + 1)]
        for i in range(n - 1, -1, -1):
            for j in range(m - 1, -1, -1):
                s = _similar(keys_a[ai[i]], keys_b[bj[j]])
                take = (s + dp[i + 1][j + 1]) if s >= threshold else -1.0
                dp[i][j] = max(take, dp[i + 1][j], dp[i][j + 1])

        out, i, j = [], 0, 0
        while i < n and j < m:
            s = _similar(keys_a[ai[i]], keys_b[bj[j]])
            if s >= threshold and abs(dp[i][j] - (s + dp[i + 1][j + 1])) < 1e-9:
                out.append((ai[i], bj[j], MOD))
                i, j = i + 1, j + 1
            elif dp[i + 1][j] >= dp[i][j + 1]:
                out.append((ai[i], None, DEL))
                i += 1
            else:
                out.append((None, bj[j], ADD))
                j += 1
        out += [(ai[k], None, DEL) for k in range(i, n)]
        out += [(None, bj[k], ADD) for k in range(j, m)]
        return out

    return pair


# ---------------------------------------------------------------- 시트 비교

def _diff_sheet(name, grid_a, grid_b, opts, keyfn, state):
    texts_a, keys_a = _prep(grid_a, keyfn)
    texts_b, keys_b = _prep(grid_b, keyfn)
    nra, nca = len(texts_a), (len(texts_a[0]) if texts_a else 0)
    nrb, ncb = len(texts_b), (len(texts_b[0]) if texts_b else 0)

    # 1) 열 정렬 — 열 전체 값을 하나의 키로
    ck_a = [tuple(keys_a[i][j] for i in range(nra)) for j in range(nca)]
    ck_b = [tuple(keys_b[i][j] for i in range(nrb)) for j in range(ncb)]
    raw_pairs = _opcode_pairs(ck_a, ck_b, _positional)

    # 행이 끼어들면 열 내용도 통째로 밀리므로, 짝지어진 열은 '같은 열'로 본다.
    # (그 열 안의 변화는 어차피 셀 단위로 표시된다)
    col_pairs = [(ja, jb, SAME if (ja is not None and jb is not None) else st)
                 for ja, jb, st in raw_pairs]

    # 2) 행 정렬 — 추가/삭제된 열을 뺀 나머지 열의 값만으로 행 키를 만든다
    paired = [(ja, jb) for ja, jb, st in col_pairs if ja is not None and jb is not None]
    base = [(ja, jb) for ja, jb in paired if _similar(ck_a[ja], ck_b[jb]) >= 0.5] or paired
    if base:
        rk_a = [tuple(keys_a[i][ja] for ja, _ in base) for i in range(nra)]
        rk_b = [tuple(keys_b[i][jb] for _, jb in base) for i in range(nrb)]
    else:
        rk_a = [tuple(r) for r in keys_a]
        rk_b = [tuple(r) for r in keys_b]

    pairer = _make_row_pairer(rk_a, rk_b, opts["similarity"])
    row_pairs = _opcode_pairs(rk_a, rk_b, pairer)

    # 3) 셀 비교
    cols = [{
        "s": st,
        "a": get_column_letter(ja + 1) if ja is not None else "",
        "b": get_column_letter(jb + 1) if jb is not None else "",
    } for ja, jb, st in col_pairs]

    stat = {"row_add": 0, "row_del": 0, "row_mod": 0,
            "col_add": 0, "col_del": 0, "cell_mod": 0}
    for _, _, st in col_pairs:
        if st == ADD:
            stat["col_add"] += 1
        elif st == DEL:
            stat["col_del"] += 1

    rows = []
    for ia, ib, st in row_pairs:
        cells = []
        if st == ADD:
            for ja, jb, cst in col_pairs:
                cells.append({"s": ADD, "n": texts_b[ib][jb]} if jb is not None else {"s": NONE})
            stat["row_add"] += 1
            rows.append({"s": ADD, "a": None, "b": ib + 1, "c": cells})
            continue
        if st == DEL:
            for ja, jb, cst in col_pairs:
                cells.append({"s": DEL, "o": texts_a[ia][ja]} if ja is not None else {"s": NONE})
            stat["row_del"] += 1
            rows.append({"s": DEL, "a": ia + 1, "b": None, "c": cells})
            continue

        changed = False
        for ja, jb, cst in col_pairs:
            if ja is None:
                cells.append({"s": ADD, "n": texts_b[ib][jb]})
            elif jb is None:
                cells.append({"s": DEL, "o": texts_a[ia][ja]})
            else:
                ka, kb = keys_a[ia][ja], keys_b[ib][jb]
                if ka == kb:
                    cells.append({"s": SAME, "n": texts_b[ib][jb]})
                else:
                    o, n = texts_a[ia][ja], texts_b[ib][jb]
                    cells.append({"s": MOD, "o": o, "n": n})
                    stat["cell_mod"] += 1
                    changed = True
                    if len(state["changes"]) < MAX_CHANGE_LIST:
                        state["changes"].append({
                            "sheet": name,
                            "a": f"{get_column_letter(ja + 1)}{ia + 1}",
                            "b": f"{get_column_letter(jb + 1)}{ib + 1}",
                            "o": o, "n": n,
                        })
                    else:
                        state["truncated"] = True
        if changed:
            stat["row_mod"] += 1
        rows.append({"s": MOD if changed else SAME, "a": ia + 1, "b": ib + 1, "c": cells})

    rows = _apply_view(rows, opts)
    total = sum(stat.values())
    return {
        "name": name,
        "status": "same" if total == 0 else "changed",
        "stat": stat,
        "cols": cols,
        "rows": rows,
        "dims": {"a": [nra, nca], "b": [nrb, ncb]},
    }


def _apply_view(rows, opts):
    """전체 보기 / 변경 주변만 보기."""
    view = opts["view"]
    if view == "all" or (view == "auto" and len(rows) <= 400):
        return rows
    ctx = max(0, int(opts["context"]))
    keep = [False] * len(rows)
    for i, r in enumerate(rows):
        if r["s"] != SAME:
            for k in range(max(0, i - ctx), min(len(rows), i + ctx + 1)):
                keep[k] = True
    if not any(keep):
        return rows[:ctx * 2] if ctx else []

    out, skipped = [], 0
    for i, r in enumerate(rows):
        if keep[i]:
            if skipped:
                out.append({"gap": skipped})
                skipped = 0
            out.append(r)
        else:
            skipped += 1
    if skipped:
        out.append({"gap": skipped})
    return out


# ---------------------------------------------------------------- 공개 API

def compare(a_bytes, b_bytes, options=None, name_a="이전", name_b="이후") -> str:
    """두 엑셀 파일을 비교하고 결과를 JSON 문자열로 돌려준다."""
    global _LAST

    opts = dict(DEFAULTS)
    if options:
        opts.update(json.loads(options) if isinstance(options, str) else dict(options))

    a_bytes, b_bytes = _to_bytes(a_bytes), _to_bytes(b_bytes)
    sig = (len(a_bytes), len(b_bytes), name_a, name_b, opts["compare"])
    if _CACHE["sig"] != sig:
        _CACHE["a"] = _parse(a_bytes, opts["compare"])
        _CACHE["b"] = _parse(b_bytes, opts["compare"])
        _CACHE["sig"] = sig
    A, B = _CACHE["a"], _CACHE["b"]

    keyfn = _mkkey(opts)
    state = {"changes": [], "truncated": False}

    # 시트 순서: 이전 파일 순서를 기본으로 하고, 새로 생긴 시트를 뒤에 붙인다
    order = list(A["order"]) + [n for n in B["order"] if n not in A["order"]]

    sheets, summary = [], {
        "sheet_add": 0, "sheet_del": 0, "sheet_mod": 0,
        "row_add": 0, "row_del": 0, "row_mod": 0,
        "col_add": 0, "col_del": 0, "cell_mod": 0,
    }
    for name in order:
        in_a, in_b = name in A["grids"], name in B["grids"]
        ga = A["grids"].get(name, [])
        gb = B["grids"].get(name, [])
        sh = _diff_sheet(name, ga, gb, opts, keyfn, state)
        if not in_b:
            sh["status"] = "deleted"
            summary["sheet_del"] += 1
        elif not in_a:
            sh["status"] = "added"
            summary["sheet_add"] += 1
        elif sh["status"] == "changed":
            summary["sheet_mod"] += 1
        for k in sh["stat"]:
            summary[k] += sh["stat"][k]
        sheets.append(sh)

    result = {
        "ok": True,
        "files": {"a": name_a, "b": name_b},
        "options": opts,
        "summary": summary,
        "sheets": sheets,
        "changes": state["changes"],
        "row_limit": MAX_ROWS,
        "row_limit_hit": any(len(g) >= MAX_ROWS
                             for src in (A["grids"], B["grids"]) for g in src.values()),
        "truncated": state["truncated"],
    }
    _LAST = result
    return json.dumps(result, ensure_ascii=False)


# ---------------------------------------------------------------- 리포트

_FILL = {
    ADD: PatternFill("solid", fgColor="DDF1E6"),
    DEL: PatternFill("solid", fgColor="FBE3E5"),
    MOD: PatternFill("solid", fgColor="FCEFD3"),
}
_LABEL = {ADD: "추가", DEL: "삭제", MOD: "수정", SAME: "동일"}


def report_xlsx(result_json: str | None = None) -> bytes:
    """비교 결과를 엑셀 리포트로 만들어 바이트로 돌려준다."""
    res = json.loads(result_json) if result_json else _LAST
    if not res:
        raise RuntimeError("비교 결과가 없습니다.")

    wb = Workbook()
    thin = Side(style="thin", color="D7DCDA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E5F")

    ws = wb.active
    ws.title = "요약"
    s = res["summary"]
    ws.append(["엑셀 변경점 리포트"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["이전 파일", res["files"]["a"]])
    ws.append(["이후 파일", res["files"]["b"]])
    ws.append(["비교 기준", "수식" if res["options"]["compare"] == "formula" else "값"])
    ws.append([])
    ws.append(["항목", "건수"])
    for c in ws[7]:
        c.font, c.fill, c.border = head, head_fill, border
    for label, key in [("시트 추가", "sheet_add"), ("시트 삭제", "sheet_del"),
                       ("시트 변경", "sheet_mod"), ("행 추가", "row_add"),
                       ("행 삭제", "row_del"), ("행 수정", "row_mod"),
                       ("열 추가", "col_add"), ("열 삭제", "col_del"),
                       ("셀 변경", "cell_mod")]:
        ws.append([label, s[key]])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 46

    ws2 = wb.create_sheet("변경 목록")
    ws2.append(["시트", "이전 위치", "이후 위치", "이전 값", "이후 값"])
    for c in ws2[1]:
        c.font, c.fill, c.border = head, head_fill, border
    for ch in res["changes"]:
        ws2.append([ch["sheet"], ch["a"], ch["b"], ch["o"], ch["n"]])
    for w, col in zip((16, 10, 10, 40, 40), "ABCDE"):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("행·열 변경")
    ws3.append(["시트", "구분", "이전 행/열", "이후 행/열", "요약"])
    for c in ws3[1]:
        c.font, c.fill, c.border = head, head_fill, border
    for sh in res["sheets"]:
        for col in sh["cols"]:
            if col["s"] in (ADD, DEL):
                ws3.append([sh["name"], f"열 {_LABEL[col['s']]}", col["a"], col["b"], ""])
                ws3.cell(ws3.max_row, 2).fill = _FILL[col["s"]]
        for row in sh["rows"]:
            if "gap" in row or row["s"] in (SAME,):
                continue
            preview = " | ".join(
                (c.get("n") or c.get("o") or "") for c in row["c"][:6]
            )[:120]
            ws3.append([sh["name"], f"행 {_LABEL[row['s']]}", row["a"], row["b"], preview])
            ws3.cell(ws3.max_row, 2).fill = _FILL[row["s"]]
    for w, col in zip((16, 10, 12, 12, 60), "ABCDE"):
        ws3.column_dimensions[col].width = w
    ws3.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_b64(result_json: str | None = None) -> str:
    """브라우저로 넘기기 쉽도록 리포트를 base64 문자열로 돌려준다."""
    import base64
    return base64.b64encode(report_xlsx(result_json)).decode("ascii")


def reset() -> None:
    """새 파일을 올릴 때 읽어둔 워크북 캐시를 비운다."""
    global _LAST
    _CACHE["sig"] = _CACHE["a"] = _CACHE["b"] = None
    _LAST = None


# ---------------------------------------------------------------- CLI

def _cli():
    import argparse

    p = argparse.ArgumentParser(description="두 엑셀 파일의 변경점을 비교합니다.")
    p.add_argument("old")
    p.add_argument("new")
    p.add_argument("--formula", action="store_true", help="값 대신 수식을 기준으로 비교")
    p.add_argument("--ignore-case", action="store_true")
    p.add_argument("--ignore-space", action="store_true")
    p.add_argument("--json", help="결과 JSON 저장 경로")
    p.add_argument("--report", help="엑셀 리포트 저장 경로")
    a = p.parse_args()

    opts = {
        "compare": "formula" if a.formula else "value",
        "ignore_case": a.ignore_case,
        "ignore_space": a.ignore_space,
        "view": "all",
    }
    with open(a.old, "rb") as f:
        ab = f.read()
    with open(a.new, "rb") as f:
        bb = f.read()

    import os
    out = compare(ab, bb, json.dumps(opts),
                  os.path.basename(a.old), os.path.basename(a.new))
    res = json.loads(out)
    s = res["summary"]

    print(f"{res['files']['a']}  →  {res['files']['b']}")
    print(f"  시트 추가 {s['sheet_add']} · 삭제 {s['sheet_del']} · 변경 {s['sheet_mod']}")
    print(f"  행  추가 {s['row_add']} · 삭제 {s['row_del']} · 수정 {s['row_mod']}")
    print(f"  열  추가 {s['col_add']} · 삭제 {s['col_del']}")
    print(f"  셀  변경 {s['cell_mod']}")
    for ch in res["changes"][:40]:
        print(f"  [{ch['sheet']}] {ch['a']} : {ch['o']!r} → {ch['n']!r}")
    if len(res["changes"]) > 40:
        print(f"  ... 외 {len(res['changes']) - 40}건")

    if a.json:
        with open(a.json, "w", encoding="utf-8") as f:
            f.write(out)
        print(f"JSON 저장: {a.json}")
    if a.report:
        with open(a.report, "wb") as f:
            f.write(report_xlsx(out))
        print(f"리포트 저장: {a.report}")


if __name__ == "__main__":
    _cli()
